"""
M M Forgings (NSE: MMFL) - independent Python replica of MMFL_Valuation_Model.xlsx

Re-computes WACC, free cash flow to firm (FCFF), terminal value and value per share,
the Bear/Base/Bull scenarios and the reverse-DCF grid, then (if the workbook is next to
this file) checks that the Excel model gives the same answers.
Run:  python mmfl_dcf.py
"""
import os

# ---------- Inputs: identical to the blue / yellow cells in the workbook ----------
CMP, SHARES = 612.0, 4.82            # Rs, crore shares (29-Sep-2026)
DEBT, INVESTMENTS = 1235.0, 10.0     # FY26, Rs crore
SALES0, NWC0 = 1590.0, 765.0         # FY26 revenue and net working capital
TAX = 0.2517                         # statutory tax rate
RF, ERP, BETA_U, KD = 0.071, 0.065, 1.0, 0.0875
G, RONIC = 0.06, 0.15                # terminal growth, return on new invested capital
GROWTH = [0.15, 0.14, 0.12, 0.10, 0.09]      # FY27E-FY31E revenue growth
MARGIN = [0.185, 0.19, 0.195, 0.195, 0.195]  # EBITDA margin
DA     = [0.068, 0.067, 0.066, 0.065, 0.064] # D&A % revenue
CAPEX  = [0.10, 0.09, 0.085, 0.085, 0.085]   # capex % revenue
NWC    = [0.46, 0.44, 0.42, 0.41, 0.40]      # NWC % revenue
FRAC   = [0.5, 1, 1, 1, 1]                   # FY27E: only the half-year after 30-Sep-2026
T      = [0.25, 1, 2, 3, 4]                  # mid-period discounting (years)
T_TV   = 4.5                                 # terminal value sits at 31-Mar-2031
SCENARIOS = {"Bear": (-0.05, -0.02, 0.04), "Base": (0, 0, 0), "Bull": (0.03, 0.02, -0.03)}  # growth, margin, NWC deltas


def wacc():
    mcap = CMP * SHARES
    beta_l = BETA_U * (1 + (1 - TAX) * DEBT / mcap)      # Hamada relevering
    ke = RF + beta_l * ERP                                # CAPM
    we = mcap / (mcap + DEBT)
    return we * ke + (1 - we) * KD * (1 - TAX), ke, beta_l


def project(growth, margin, dg=0.0, dm=0.0, dn=0.0):
    """Five forecast years. Returns a list of dicts with sales, NOPAT and FCFF."""
    out, s_prev, nwc_prev = [], SALES0, NWC0
    for i in range(5):
        s = s_prev * (1 + growth[i] + dg)
        ebitda = s * (margin[i] + dm)
        da = s * DA[i]
        nopat = (ebitda - da) * (1 - TAX)
        nwc = s * (NWC[i] + dn)
        fcff = nopat + da - s * CAPEX[i] - (nwc - nwc_prev)
        out.append({"sales": s, "ebitda": ebitda, "nopat": nopat, "fcff": fcff})
        s_prev, nwc_prev = s, nwc
    return out


def value_per_share(years, w, g=G, ronic=RONIC):
    pv_fcff = sum(y["fcff"] * f / (1 + w) ** t for y, f, t in zip(years, FRAC, T))
    tv = years[-1]["nopat"] * (1 + g) * (1 - g / ronic) / (w - g)   # value-driver formula
    ev = pv_fcff + tv / (1 + w) ** T_TV
    return (ev - DEBT + INVESTMENTS) / SHARES


def grid(w, cagrs, margins):
    return [[value_per_share(project([c] * 5, [m] * 5), w) for m in margins] for c in cagrs]


if __name__ == "__main__":
    w, ke, beta_l = wacc()
    print(f"Levered beta {beta_l:.2f} | cost of equity {ke:.2%} | WACC {w:.2%}")
    res = {}
    for name, (dg, dm, dn) in SCENARIOS.items():
        res[name] = value_per_share(project(GROWTH, MARGIN, dg, dm, dn), w)
        print(f"{name:5s} value per share Rs{res[name]:7.1f}  ({res[name] / CMP - 1:+.0%} vs CMP)")
    cagrs = [0.06, 0.08, 0.10, 0.12, 0.14, 0.16, 0.18, 0.20, 0.22]
    margins = [0.16, 0.18, 0.20, 0.22, 0.24, 0.26]
    grids = {"A": grid(w, cagrs, margins), "B": grid(0.11, cagrs, margins)}
    for key, label in [("A", f"WACC {w:.2%}"), ("B", "WACC 11.00%")]:
        print(f"\nGrid {key} ({label}): value per share, rows = revenue CAGR, cols = EBITDA margin")
        print("       " + "".join(f"{m:>7.0%}" for m in margins))
        for c, rowv in zip(cagrs, grids[key]):
            print(f"{c:>6.0%} " + "".join(f"{v:7.0f}" for v in rowv))
    # ---- cross-check against the Excel workbook (needs openpyxl and a recalculated file) ----
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "MMFL_Valuation_Model.xlsx")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        xl_base = wb["DCF"]["B49"].value
        g_ws = wb["Reverse_DCF_Grid"]
        diffs = [abs(g_ws.cell(25 + i, 3 + j).value - grids["A"][i][j]) for i in range(9) for j in range(6)]
        diffs += [abs(g_ws.cell(37 + i, 3 + j).value - grids["B"][i][j]) for i in range(9) for j in range(6)]
        print(f"\nExcel base value Rs{xl_base:.2f} vs Python Rs{res['Base']:.2f} | max grid difference Rs{max(diffs):.4f}")
    except Exception as e:  # file missing or not recalculated
        print(f"\n(Excel cross-check skipped: {e})")
